import io
import os
import sqlite3
import secrets as secrets_module
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

app = Flask(__name__)
# In production (Render), set SECRET_KEY as an environment variable.
# Locally, this generates a temporary one each run, which is fine for testing.
app.secret_key = os.environ.get('SECRET_KEY') or secrets_module.token_hex(32)

# --- SCHOOL / ELECTION SETTINGS ---
SCHOOL_NAME = "GREAT AUNTY AYO COLLEGE, MOWE"
ACADEMIC_SESSION = "2026/2027"
ELECTION_YEAR = 2026
# In production (Render), set ADMIN_PASSWORD as an environment variable.
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'changeme2026')

PREFECT_POSITIONS = [
    "Head Boy",
    "Head Girl",
    "Asst. Head Boy",
    "Asst. Head Girl",
    "Health Prefect",
    "Sports Prefect",
    "Library Prefect",
    "Social Prefect",
    "Chapel Prefect",
]

# --- DATABASE ENGINE ---
# If DATABASE_URL is set (e.g. on Render, pointing at a Postgres instance),
# use Postgres. Otherwise fall back to a local SQLite file for easy local testing.
DATABASE_URL = os.environ.get('DATABASE_URL')
USE_POSTGRES = bool(DATABASE_URL)

if USE_POSTGRES:
    import psycopg2
    import psycopg2.extras


@app.context_processor
def inject_school_info():
    """Makes these available in every template automatically."""
    return {
        "school_name": SCHOOL_NAME,
        "academic_session": ACADEMIC_SESSION,
        "election_year": ELECTION_YEAR,
    }


class _PGCursorWrapper:
    """Makes a psycopg2 cursor behave like a sqlite3 cursor for our purposes."""
    def __init__(self, cursor):
        self._cursor = cursor

    def fetchone(self):
        return self._cursor.fetchone()

    def fetchall(self):
        return self._cursor.fetchall()


class _PGConnWrapper:
    """Wraps a psycopg2 connection so route code can keep calling
    conn.execute(query, params) the same way it does with sqlite3, using
    '?' placeholders. This means the rest of the app never has to know
    which database engine is actually running underneath."""
    def __init__(self, conn):
        self._conn = conn

    def execute(self, query, params=()):
        cur = self._conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute(query.replace('?', '%s'), params)
        return _PGCursorWrapper(cur)

    def commit(self):
        self._conn.commit()

    def close(self):
        self._conn.close()


def get_db_connection():
    if USE_POSTGRES:
        raw_conn = psycopg2.connect(DATABASE_URL, sslmode='require')
        return _PGConnWrapper(raw_conn)
    else:
        conn = sqlite3.connect('database.db')
        conn.row_factory = sqlite3.Row
        return conn


# Initialize Database Tables
def init_db():
    conn = get_db_connection()

    if USE_POSTGRES:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS users (
                voter_id TEXT PRIMARY KEY,
                full_name TEXT NOT NULL,
                role TEXT CHECK(role IN ('Student', 'Teacher', 'Admin')),
                has_voted INTEGER DEFAULT 0,
                password_hash TEXT
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS candidates (
                id SERIAL PRIMARY KEY,
                full_name TEXT NOT NULL,
                position TEXT NOT NULL,
                manifesto TEXT NOT NULL,
                is_approved INTEGER DEFAULT 0,
                vote_count INTEGER DEFAULT 0
            )
        ''')
        # Safe to run every startup: no-op if the column already exists.
        conn.execute('ALTER TABLE users ADD COLUMN IF NOT EXISTS password_hash TEXT')
    else:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS users (
                voter_id TEXT PRIMARY KEY,
                full_name TEXT NOT NULL,
                role TEXT CHECK(role IN ('Student', 'Teacher', 'Admin')),
                has_voted INTEGER DEFAULT 0,
                password_hash TEXT
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS candidates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                full_name TEXT NOT NULL,
                position TEXT NOT NULL,
                manifesto TEXT NOT NULL,
                is_approved INTEGER DEFAULT 0,
                vote_count INTEGER DEFAULT 0
            )
        ''')
        existing_columns = [row['name'] for row in conn.execute('PRAGMA table_info(users)').fetchall()]
        if 'password_hash' not in existing_columns:
            conn.execute('ALTER TABLE users ADD COLUMN password_hash TEXT')

    conn.commit()
    conn.close()


init_db()


# --- ACCESS CONTROL ---
def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('is_admin'):
            flash("Please log in as admin to continue.", "error")
            return redirect(url_for('admin_login'))
        return f(*args, **kwargs)
    return decorated


# --- ROUTES ---

@app.route('/')
def index():
    return redirect(url_for('login'))


# 1. Voter Login
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        voter_id = request.form['voter_id'].strip().upper()
        password = request.form.get('password', '')
        conn = get_db_connection()
        user = conn.execute('SELECT * FROM users WHERE voter_id = ?', (voter_id,)).fetchone()
        conn.close()

        if user is None:
            flash("Invalid Voter ID. Please check with the school administrator.", "error")
        elif not user['password_hash']:
            # Account exists but has no password set yet (e.g. an old record).
            flash("Your account has no password set. Please see the admin to set one up.", "error")
        elif not check_password_hash(user['password_hash'], password):
            flash("Incorrect password.", "error")
        else:
            session['voter_id'] = user['voter_id']
            session['role'] = user['role']
            session['has_voted'] = user['has_voted']
            flash(f"Welcome, {user['full_name']}!", "success")
            return redirect(url_for('vote'))

    return render_template('login.html')


# 1b. Change Password (for a logged-in voter)
@app.route('/change-password', methods=['GET', 'POST'])
def change_password():
    if 'voter_id' not in session:
        return redirect(url_for('login'))

    if request.method == 'POST':
        current_password = request.form.get('current_password', '')
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')

        conn = get_db_connection()
        user = conn.execute('SELECT * FROM users WHERE voter_id = ?', (session['voter_id'],)).fetchone()

        if not user or not check_password_hash(user['password_hash'], current_password):
            conn.close()
            flash("Current password is incorrect.", "error")
        elif len(new_password) < 6:
            conn.close()
            flash("New password must be at least 6 characters.", "error")
        elif new_password != confirm_password:
            conn.close()
            flash("New passwords do not match.", "error")
        else:
            conn.execute(
                'UPDATE users SET password_hash = ? WHERE voter_id = ?',
                (generate_password_hash(new_password), session['voter_id'])
            )
            conn.commit()
            conn.close()
            flash("Password updated successfully.", "success")
            return redirect(url_for('vote'))

    return render_template('change_password.html')


# 2. Prefect Candidate Registration Form
@app.route('/register-candidate', methods=['GET', 'POST'])
def register_candidate():
    if request.method == 'POST':
        full_name = request.form['full_name'].strip()
        position = request.form['position']
        manifesto = request.form['manifesto'].strip()

        if position not in PREFECT_POSITIONS:
            flash("Invalid position selected.", "error")
            return redirect(url_for('register_candidate'))

        conn = get_db_connection()
        conn.execute(
            'INSERT INTO candidates (full_name, position, manifesto) VALUES (?, ?, ?)',
            (full_name, position, manifesto)
        )
        conn.commit()
        conn.close()

        flash("Registration submitted! Pending admin approval.", "info")
        return redirect(url_for('register_candidate'))

    return render_template('register.html', positions=PREFECT_POSITIONS)


# 3. Voting Interface
@app.route('/vote', methods=['GET', 'POST'])
def vote():
    if 'voter_id' not in session:
        return redirect(url_for('login'))

    conn = get_db_connection()
    user = conn.execute('SELECT * FROM users WHERE voter_id = ?', (session['voter_id'],)).fetchone()

    if user is None:
        conn.close()
        session.clear()
        return redirect(url_for('login'))

    if user['has_voted'] == 1:
        conn.close()
        return render_template('already_voted.html')

    if request.method == 'POST':
        submitted = request.form.to_dict()

        # Validate every submitted candidate actually belongs to the position
        # it was submitted under, so a forged/edited form can't stuff votes.
        for position, candidate_id in submitted.items():
            if position not in PREFECT_POSITIONS:
                continue
            candidate = conn.execute(
                'SELECT * FROM candidates WHERE id = ? AND position = ? AND is_approved = 1',
                (candidate_id, position)
            ).fetchone()
            if candidate is None:
                conn.close()
                flash("Invalid ballot submission. Please try again.", "error")
                return redirect(url_for('vote'))

        for position, candidate_id in submitted.items():
            if position not in PREFECT_POSITIONS:
                continue
            conn.execute('UPDATE candidates SET vote_count = vote_count + 1 WHERE id = ?', (candidate_id,))

        conn.execute('UPDATE users SET has_voted = 1 WHERE voter_id = ?', (session['voter_id'],))
        conn.commit()
        conn.close()

        session['has_voted'] = 1
        flash("Your votes have been recorded successfully!", "success")
        return redirect(url_for('vote'))

    approved_candidates = conn.execute(
        'SELECT * FROM candidates WHERE is_approved = 1 ORDER BY position, full_name'
    ).fetchall()
    conn.close()

    positions = {}
    for candidate in approved_candidates:
        pos = candidate['position']
        positions.setdefault(pos, []).append(candidate)

    return render_template('vote.html', positions=positions, voter_name=user['full_name'])


# 4. Admin Login
@app.route('/admin-login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        password = request.form.get('password', '')
        if password == ADMIN_PASSWORD:
            session['is_admin'] = True
            flash("Admin login successful.", "success")
            return redirect(url_for('admin'))
        else:
            flash("Incorrect admin password.", "error")
    return render_template('admin_login.html')


@app.route('/admin-logout')
def admin_logout():
    session.pop('is_admin', None)
    flash("Admin logged out.", "info")
    return redirect(url_for('login'))


# 5. Admin Dashboard (Approve candidates & see live results)
@app.route('/admin', methods=['GET', 'POST'])
@admin_required
def admin():
    conn = get_db_connection()

    if request.method == 'POST':
        candidate_id = request.form.get('candidate_id')
        action = request.form.get('action')

        if action == 'approve':
            conn.execute('UPDATE candidates SET is_approved = 1 WHERE id = ?', (candidate_id,))
        elif action == 'reject':
            conn.execute('DELETE FROM candidates WHERE id = ?', (candidate_id,))
        conn.commit()

    pending_candidates = conn.execute('SELECT * FROM candidates WHERE is_approved = 0').fetchall()
    all_candidates = conn.execute(
        'SELECT * FROM candidates WHERE is_approved = 1 ORDER BY position, vote_count DESC'
    ).fetchall()

    total_voters = conn.execute('SELECT COUNT(*) AS c FROM users').fetchone()['c']
    total_voted = conn.execute('SELECT COUNT(*) AS c FROM users WHERE has_voted = 1').fetchone()['c']
    conn.close()

    return render_template(
        'admin.html',
        pending=pending_candidates,
        candidates=all_candidates,
        total_voters=total_voters,
        total_voted=total_voted,
    )


# 6. Add Voters (replaces running a local script — everything happens on the website)
def _insert_voter(conn, voter_id, full_name, role):
    """Inserts one voter if valid and not already present.
    Returns a tuple: ('added' | 'skipped' | 'error', message)."""
    voter_id = (voter_id or '').strip().upper()
    full_name = (full_name or '').strip()
    role = (role or '').strip()

    if not voter_id or not full_name or not role:
        return ('error', f"Incomplete row — ID: \"{voter_id}\", Name: \"{full_name}\", Role: \"{role}\"")
    if role not in ('Student', 'Teacher', 'Admin'):
        return ('error', f"{voter_id}: role must be Student, Teacher, or Admin — got \"{role}\"")

    existing = conn.execute('SELECT voter_id FROM users WHERE voter_id = ?', (voter_id,)).fetchone()
    if existing:
        return ('skipped', voter_id)

    conn.execute(
        'INSERT INTO users (voter_id, full_name, role, password_hash) VALUES (?, ?, ?, ?)',
        (voter_id, full_name, role, generate_password_hash(voter_id))
    )
    return ('added', voter_id)


@app.route('/admin/add-voters', methods=['GET', 'POST'])
@admin_required
def add_voters():
    results = None
    if request.method == 'POST':
        added, skipped, errors = [], [], []
        conn = get_db_connection()

        uploaded_file = request.files.get('voter_file')
        if uploaded_file and uploaded_file.filename:
            if not uploaded_file.filename.lower().endswith('.xlsx'):
                errors.append("Please upload a .xlsx file (the template file works best).")
            else:
                try:
                    wb = load_workbook(uploaded_file, data_only=True)
                    ws = wb.active
                    # Row 1 is treated as the header and skipped.
                    for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
                        voter_id, full_name, role = (row + (None, None, None))[:3]
                        if not voter_id and not full_name and not role:
                            continue  # skip fully blank rows
                        status, message = _insert_voter(conn, voter_id, full_name, role)
                        if status == 'added':
                            added.append(message)
                        elif status == 'skipped':
                            skipped.append(message)
                        else:
                            errors.append(message)
                except Exception:
                    errors.append("Could not read that file. Make sure it's a valid .xlsx file, ideally the downloaded template.")
        else:
            raw_text = request.form.get('voter_list', '')
            for line_number, line in enumerate(raw_text.splitlines(), start=1):
                line = line.strip()
                if not line:
                    continue
                parts = [p.strip() for p in line.split(',')]
                if len(parts) != 3:
                    errors.append(f"Line {line_number}: expected 'ID, Full Name, Role' — got \"{line}\"")
                    continue
                status, message = _insert_voter(conn, *parts)
                if status == 'added':
                    added.append(message)
                elif status == 'skipped':
                    skipped.append(message)
                else:
                    errors.append(message)

        conn.commit()
        conn.close()
        results = {"added": added, "skipped": skipped, "errors": errors}

    return render_template('add_voters.html', results=results)


# 6b. Downloadable Excel template for bulk-adding voters
@app.route('/admin/voter-template')
@admin_required
def voter_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Voters"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    example_font = Font(italic=True, color="6B7280")
    instructions_font = Font(italic=True, size=10, color="B45309")

    headers = ["Voter ID", "Full Name", "Role"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    example_rows = [
        ("STU101", "John Doe", "Student"),
        ("TCH001", "Mr. Smith", "Teacher"),
    ]
    for row in example_rows:
        ws.append(row)
    for row_num in range(2, 2 + len(example_rows)):
        for col in range(1, 4):
            ws.cell(row=row_num, column=col).font = example_font

    note_row = 2 + len(example_rows) + 1
    ws.cell(row=note_row, column=1,
             value="↑ Replace the rows above with your real voters, then upload this file on the Add Voters page.").font = instructions_font

    # Dropdown restricting the Role column to valid values.
    role_validation = DataValidation(type="list", formula1='"Student,Teacher,Admin"', allow_blank=True)
    ws.add_data_validation(role_validation)
    role_validation.add(f"C2:C500")

    for col, width in zip('ABC', (16, 28, 14)):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="voter_upload_template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# 7. Export Live Results to Excel
@app.route('/admin/export-results')
@admin_required
def export_results():
    conn = get_db_connection()
    candidates = conn.execute(
        'SELECT * FROM candidates WHERE is_approved = 1 ORDER BY position, vote_count DESC'
    ).fetchall()
    total_voters = conn.execute('SELECT COUNT(*) AS c FROM users').fetchone()['c']
    total_voted = conn.execute('SELECT COUNT(*) AS c FROM users WHERE has_voted = 1').fetchone()['c']
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Election Results"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")

    ws.append([f"{SCHOOL_NAME} — Prefectship Election {ELECTION_YEAR}"])
    ws.append([f"Academic Session: {ACADEMIC_SESSION}"])
    ws.append([f"Voter Turnout: {total_voted} of {total_voters} voters"])
    ws.append([])
    ws.append(["Position", "Candidate Name", "Total Votes", "Result"])

    for cell in ws[5]:
        cell.font = header_font
        cell.fill = header_fill

    current_position = None
    for candidate in candidates:
        is_leading = current_position != candidate['position']
        current_position = candidate['position']
        ws.append([
            candidate['position'],
            candidate['full_name'],
            candidate['vote_count'],
            "Leading" if is_leading else "",
        ])

    for col, width in zip('ABCD', (22, 28, 14, 12)):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"election_results_{ELECTION_YEAR}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


if __name__ == '__main__':
    app.run(debug=True) use Postgres. Otherwise fall back to a local SQLite file for easy local testing.
DATABASE_URL = os.environ.get('DATABASE_URL')
USE_POSTGRES = bool(DATABASE_URL)

if USE_POSTGRES:
    import psycopg2
    import psycopg2.extras


@app.context_processor
def inject_school_info():
    """Makes these available in every template automatically."""
    return {
        "school_name": SCHOOL_NAME,
        "academic_session": ACADEMIC_SESSION,
        "election_year": ELECTION_YEAR,
    }


class _PGCursorWrapper:
    """Makes a psycopg2 cursor behave like a sqlite3 cursor for our purposes."""
    def __init__(self, cursor):
        self._cursor = cursor

    def fetchone(self):
        return self._cursor.fetchone()

    def fetchall(self):
        return self._cursor.fetchall()


class _PGConnWrapper:
    """Wraps a psycopg2 connection so route code can keep calling
    conn.execute(query, params) the same way it does with sqlite3, using
    '?' placeholders. This means the rest of the app never has to know
    which database engine is actually running underneath."""
    def __init__(self, conn):
        self._conn = conn

    def execute(self, query, params=()):
        cur = self._conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute(query.replace('?', '%s'), params)
        return _PGCursorWrapper(cur)

    def commit(self):
        self._conn.commit()

    def close(self):
        self._conn.close()


def get_db_connection():
    if USE_POSTGRES:
        raw_conn = psycopg2.connect(DATABASE_URL, sslmode='require')
        return _PGConnWrapper(raw_conn)
    else:
        conn = sqlite3.connect('database.db')
        conn.row_factory = sqlite3.Row
        return conn


# Initialize Database Tables
def init_db():
    conn = get_db_connection()

    if USE_POSTGRES:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS users (
                voter_id TEXT PRIMARY KEY,
                full_name TEXT NOT NULL,
                role TEXT CHECK(role IN ('Student', 'Teacher', 'Admin')),
                has_voted INTEGER DEFAULT 0,
                password_hash TEXT
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS candidates (
                id SERIAL PRIMARY KEY,
                full_name TEXT NOT NULL,
                position TEXT NOT NULL,
                manifesto TEXT NOT NULL,
                is_approved INTEGER DEFAULT 0,
                vote_count INTEGER DEFAULT 0
            )
        ''')
        # Safe to run every startup: no-op if the column already exists.
        conn.execute('ALTER TABLE users ADD COLUMN IF NOT EXISTS password_hash TEXT')
    else:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS users (
                voter_id TEXT PRIMARY KEY,
                full_name TEXT NOT NULL,
                role TEXT CHECK(role IN ('Student', 'Teacher', 'Admin')),
                has_voted INTEGER DEFAULT 0,
                password_hash TEXT
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS candidates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                full_name TEXT NOT NULL,
                position TEXT NOT NULL,
                manifesto TEXT NOT NULL,
                is_approved INTEGER DEFAULT 0,
                vote_count INTEGER DEFAULT 0
            )
        ''')
        existing_columns = [row['name'] for row in conn.execute('PRAGMA table_info(users)').fetchall()]
        if 'password_hash' not in existing_columns:
            conn.execute('ALTER TABLE users ADD COLUMN password_hash TEXT')

    conn.commit()
    conn.close()


init_db()


# --- ACCESS CONTROL ---
def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('is_admin'):
            flash("Please log in as admin to continue.", "error")
            return redirect(url_for('admin_login'))
        return f(*args, **kwargs)
    return decorated


# --- ROUTES ---

@app.route('/')
def index():
    return redirect(url_for('login'))


# 1. Voter Login
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        voter_id = request.form['voter_id'].strip().upper()
        password = request.form.get('password', '')
        conn = get_db_connection()
        user = conn.execute('SELECT * FROM users WHERE voter_id = ?', (voter_id,)).fetchone()
        conn.close()

        if user is None:
            flash("Invalid Voter ID. Please check with the school administrator.", "error")
        elif not user['password_hash']:
            # Account exists but has no password set yet (e.g. an old record).
            flash("Your account has no password set. Please see the admin to set one up.", "error")
        elif not check_password_hash(user['password_hash'], password):
            flash("Incorrect password.", "error")
        else:
            session['voter_id'] = user['voter_id']
            session['role'] = user['role']
            session['has_voted'] = user['has_voted']
            flash(f"Welcome, {user['full_name']}!", "success")
            return redirect(url_for('vote'))

    return render_template('login.html')


# 1b. Change Password (for a logged-in voter)
@app.route('/change-password', methods=['GET', 'POST'])
def change_password():
    if 'voter_id' not in session:
        return redirect(url_for('login'))

    if request.method == 'POST':
        current_password = request.form.get('current_password', '')
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')

        conn = get_db_connection()
        user = conn.execute('SELECT * FROM users WHERE voter_id = ?', (session['voter_id'],)).fetchone()

        if not user or not check_password_hash(user['password_hash'], current_password):
            conn.close()
            flash("Current password is incorrect.", "error")
        elif len(new_password) < 6:
            conn.close()
            flash("New password must be at least 6 characters.", "error")
        elif new_password != confirm_password:
            conn.close()
            flash("New passwords do not match.", "error")
        else:
            conn.execute(
                'UPDATE users SET password_hash = ? WHERE voter_id = ?',
                (generate_password_hash(new_password), session['voter_id'])
            )
            conn.commit()
            conn.close()
            flash("Password updated successfully.", "success")
            return redirect(url_for('vote'))

    return render_template('change_password.html')


# 2. Prefect Candidate Registration Form
@app.route('/register-candidate', methods=['GET', 'POST'])
def register_candidate():
    if request.method == 'POST':
        full_name = request.form['full_name'].strip()
        position = request.form['position']
        manifesto = request.form['manifesto'].strip()

        if position not in PREFECT_POSITIONS:
            flash("Invalid position selected.", "error")
            return redirect(url_for('register_candidate'))

        conn = get_db_connection()
        conn.execute(
            'INSERT INTO candidates (full_name, position, manifesto) VALUES (?, ?, ?)',
            (full_name, position, manifesto)
        )
        conn.commit()
        conn.close()

        flash("Registration submitted! Pending admin approval.", "info")
        return redirect(url_for('register_candidate'))

    return render_template('register.html', positions=PREFECT_POSITIONS)


# 3. Voting Interface
@app.route('/vote', methods=['GET', 'POST'])
def vote():
    if 'voter_id' not in session:
        return redirect(url_for('login'))

    conn = get_db_connection()
    user = conn.execute('SELECT * FROM users WHERE voter_id = ?', (session['voter_id'],)).fetchone()

    if user is None:
        conn.close()
        session.clear()
        return redirect(url_for('login'))

    if user['has_voted'] == 1:
        conn.close()
        return render_template('already_voted.html')

    if request.method == 'POST':
        submitted = request.form.to_dict()

        # Validate every submitted candidate actually belongs to the position
        # it was submitted under, so a forged/edited form can't stuff votes.
        for position, candidate_id in submitted.items():
            if position not in PREFECT_POSITIONS:
                continue
            candidate = conn.execute(
                'SELECT * FROM candidates WHERE id = ? AND position = ? AND is_approved = 1',
                (candidate_id, position)
            ).fetchone()
            if candidate is None:
                conn.close()
                flash("Invalid ballot submission. Please try again.", "error")
                return redirect(url_for('vote'))

        for position, candidate_id in submitted.items():
            if position not in PREFECT_POSITIONS:
                continue
            conn.execute('UPDATE candidates SET vote_count = vote_count + 1 WHERE id = ?', (candidate_id,))

        conn.execute('UPDATE users SET has_voted = 1 WHERE voter_id = ?', (session['voter_id'],))
        conn.commit()
        conn.close()

        session['has_voted'] = 1
        flash("Your votes have been recorded successfully!", "success")
        return redirect(url_for('vote'))

    approved_candidates = conn.execute(
        'SELECT * FROM candidates WHERE is_approved = 1 ORDER BY position, full_name'
    ).fetchall()
    conn.close()

    positions = {}
    for candidate in approved_candidates:
        pos = candidate['position']
        positions.setdefault(pos, []).append(candidate)

    return render_template('vote.html', positions=positions, voter_name=user['full_name'])


# 4. Admin Login
@app.route('/admin-login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        password = request.form.get('password', '')
        if password == ADMIN_PASSWORD:
            session['is_admin'] = True
            flash("Admin login successful.", "success")
            return redirect(url_for('admin'))
        else:
            flash("Incorrect admin password.", "error")
    return render_template('admin_login.html')


@app.route('/admin-logout')
def admin_logout():
    session.pop('is_admin', None)
    flash("Admin logged out.", "info")
    return redirect(url_for('login'))


# 5. Admin Dashboard (Approve candidates & see live results)
@app.route('/admin', methods=['GET', 'POST'])
@admin_required
def admin():
    conn = get_db_connection()

    if request.method == 'POST':
        candidate_id = request.form.get('candidate_id')
        action = request.form.get('action')

        if action == 'approve':
            conn.execute('UPDATE candidates SET is_approved = 1 WHERE id = ?', (candidate_id,))
        elif action == 'reject':
            conn.execute('DELETE FROM candidates WHERE id = ?', (candidate_id,))
        conn.commit()

    pending_candidates = conn.execute('SELECT * FROM candidates WHERE is_approved = 0').fetchall()
    all_candidates = conn.execute(
        'SELECT * FROM candidates WHERE is_approved = 1 ORDER BY position, vote_count DESC'
    ).fetchall()

    total_voters = conn.execute('SELECT COUNT(*) AS c FROM users').fetchone()['c']
    total_voted = conn.execute('SELECT COUNT(*) AS c FROM users WHERE has_voted = 1').fetchone()['c']
    conn.close()

    return render_template(
        'admin.html',
        pending=pending_candidates,
        candidates=all_candidates,
        total_voters=total_voters,
        total_voted=total_voted,
    )


# 6. Add Voters (replaces running a local script — everything happens on the website)
@app.route('/admin/add-voters', methods=['GET', 'POST'])
@admin_required
def add_voters():
    results = None
    if request.method == 'POST':
        raw_text = request.form.get('voter_list', '')
        added, skipped, errors = [], [], []

        conn = get_db_connection()
        for line_number, line in enumerate(raw_text.splitlines(), start=1):
            line = line.strip()
            if not line:
                continue
            parts = [p.strip() for p in line.split(',')]
            if len(parts) != 3:
                errors.append(f"Line {line_number}: expected 'ID, Full Name, Role' — got \"{line}\"")
                continue

            voter_id, full_name, role = parts
            voter_id = voter_id.upper()
            if role not in ('Student', 'Teacher', 'Admin'):
                errors.append(f"Line {line_number}: role must be Student, Teacher, or Admin — got \"{role}\"")
                continue

            existing = conn.execute('SELECT voter_id FROM users WHERE voter_id = ?', (voter_id,)).fetchone()
            if existing:
                skipped.append(voter_id)
                continue

            conn.execute(
                'INSERT INTO users (voter_id, full_name, role, password_hash) VALUES (?, ?, ?, ?)',
                (voter_id, full_name, role, generate_password_hash(voter_id))
            )
            added.append(voter_id)

        conn.commit()
        conn.close()
        results = {"added": added, "skipped": skipped, "errors": errors}

    return render_template('add_voters.html', results=results)


# 7. Export Live Results to Excel
@app.route('/admin/export-results')
@admin_required
def export_results():
    conn = get_db_connection()
    candidates = conn.execute(
        'SELECT * FROM candidates WHERE is_approved = 1 ORDER BY position, vote_count DESC'
    ).fetchall()
    total_voters = conn.execute('SELECT COUNT(*) AS c FROM users').fetchone()['c']
    total_voted = conn.execute('SELECT COUNT(*) AS c FROM users WHERE has_voted = 1').fetchone()['c']
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Election Results"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")

    ws.append([f"{SCHOOL_NAME} — Prefectship Election {ELECTION_YEAR}"])
    ws.append([f"Academic Session: {ACADEMIC_SESSION}"])
    ws.append([f"Voter Turnout: {total_voted} of {total_voters} voters"])
    ws.append([])
    ws.append(["Position", "Candidate Name", "Total Votes", "Result"])

    for cell in ws[5]:
        cell.font = header_font
        cell.fill = header_fill

    current_position = None
    for candidate in candidates:
        is_leading = current_position != candidate['position']
        current_position = candidate['position']
        ws.append([
            candidate['position'],
            candidate['full_name'],
            candidate['vote_count'],
            "Leading" if is_leading else "",
        ])

    for col, width in zip('ABCD', (22, 28, 14, 12)):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"election_results_{ELECTION_YEAR}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


if __name__ == '__main__':
    app.run(debug=True)
